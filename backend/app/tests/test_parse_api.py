from __future__ import annotations

import logging
from pathlib import Path

import fitz
import pytest
from docx import Document as DocxDocument
from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy.orm import Session

from backend.app.core.config import settings
from backend.app.db.models import Document, DocumentParseResult
from backend.app.db.session import SessionLocal
from backend.app.main import app
from backend.app.services import parse_service


def _upload(name: str, content: bytes, content_type: str = "text/plain") -> dict:
    with TestClient(app) as client:
        response = client.post("/api/upload", files={"file": (name, content, content_type)})
    assert response.status_code == 200
    return response.json()["data"]


def _parse(document_id: str):
    with TestClient(app) as client:
        return client.post(f"/api/documents/{document_id}/parse")


def _document(document_id: str) -> Document:
    with SessionLocal() as db:
        document = db.get(Document, document_id)
        assert document is not None
        return document


def test_txt_parse_success_writes_parsed_file() -> None:
    document = _upload("parse-me.txt", "hello parse".encode("utf-8"))

    response = _parse(document["id"])

    assert response.status_code == 200
    data = response.json()["data"]
    assert set(data.keys()) == {
        "document_id",
        "status",
        "source_sha256",
        "parsed_relative_path",
        "parser_name",
        "parser_version",
        "content_sha256",
        "char_count",
        "truncated",
        "error_message",
        "parsed_at",
    }
    assert data["status"] == "parsed"
    assert data["source_sha256"] == document["sha256"]
    assert data["parsed_relative_path"].startswith("parsed/")
    parsed_path = settings.data_path / data["parsed_relative_path"]
    assert parsed_path.exists()
    assert parsed_path.read_text(encoding="utf-8") == "hello parse"


def test_markdown_parse_reuses_existing_result() -> None:
    document = _upload("reuse.md", b"# Reuse")

    first = _parse(document["id"])
    second = _parse(document["id"])

    assert first.status_code == 200
    assert second.status_code == 200
    assert first.json()["data"]["parsed_at"] == second.json()["data"]["parsed_at"]
    assert first.json()["data"]["parsed_relative_path"] == second.json()["data"]["parsed_relative_path"]


def test_archived_document_parse_returns_400() -> None:
    document = _upload("archived.md", b"archived")
    with TestClient(app) as client:
        client.patch(f"/api/documents/{document['id']}/status", json={"status": "archived"})
        response = client.post(f"/api/documents/{document['id']}/parse")

    assert response.status_code == 400


def test_sha256_mismatch_rejects_parse() -> None:
    document = _upload("changed.md", b"original")
    stored = _document(document["id"])
    (settings.data_path / stored.relative_path).write_bytes(b"changed")

    response = _parse(document["id"])

    assert response.status_code == 400
    assert "checksum" in response.json()["detail"]["message"].lower()


def test_failed_parse_can_retry(monkeypatch: pytest.MonkeyPatch) -> None:
    document = _upload("retry.md", b"retry")

    def fail_once(path: Path, extension: str):
        raise RuntimeError("temporary parser failure")

    monkeypatch.setattr(parse_service, "parse_file", fail_once)
    failed = _parse(document["id"])
    assert failed.status_code == 400

    monkeypatch.undo()
    retried = _parse(document["id"])
    assert retried.status_code == 200
    assert retried.json()["data"]["status"] == "parsed"


def test_missing_file_records_failed_without_dirty_file() -> None:
    document = _upload("missing.md", b"missing")
    stored = _document(document["id"])
    (settings.data_path / stored.relative_path).unlink()

    response = _parse(document["id"])

    assert response.status_code == 400
    with SessionLocal() as db:
        result = db.get(DocumentParseResult, document["id"])
        assert result is not None
        assert result.status == "failed"
        assert result.parsed_relative_path is None
    assert not list(settings.parsed_path.rglob(f"{document['id']}*.tmp*"))


def test_db_failure_cleans_generated_parsed_file(monkeypatch: pytest.MonkeyPatch) -> None:
    document = _upload("db-failure.md", b"db failure")

    def fail_commit(self: Session) -> None:
        raise RuntimeError("forced parse commit failure")

    monkeypatch.setattr(Session, "commit", fail_commit)
    with TestClient(app, raise_server_exceptions=False) as client:
        response = client.post(f"/api/documents/{document['id']}/parse")

    assert response.status_code == 500
    assert not list(settings.parsed_path.rglob(f"{document['id']}.txt"))


def test_delete_document_removes_parsed_file_and_parse_result() -> None:
    document = _upload("delete-parsed.md", b"delete parsed")
    parsed = _parse(document["id"]).json()["data"]
    parsed_path = settings.data_path / parsed["parsed_relative_path"]
    assert parsed_path.exists()

    with TestClient(app) as client:
        response = client.delete(f"/api/documents/{document['id']}")

    assert response.status_code == 200
    assert not parsed_path.exists()
    with SessionLocal() as db:
        assert db.get(DocumentParseResult, document["id"]) is None
        assert db.get(Document, document["id"]) is None


def test_delete_refuses_unsafe_parsed_path_and_logs_error(caplog) -> None:
    document = _upload("unsafe-parsed.md", b"unsafe parsed")
    _parse(document["id"])
    with SessionLocal() as db:
        result = db.get(DocumentParseResult, document["id"])
        assert result is not None
        result.parsed_relative_path = "../sqlite/workmemory.db"
        db.commit()

    caplog.set_level(logging.ERROR)
    with TestClient(app) as client:
        response = client.delete(f"/api/documents/{document['id']}")

    assert response.status_code == 500
    assert response.json()["detail"]["code"] == "UNSAFE_PARSED_FILE_PATH"
    assert "outside data/parsed" in caplog.text
    with SessionLocal() as db:
        assert db.get(Document, document["id"]) is not None
        assert db.get(DocumentParseResult, document["id"]) is not None


def test_csv_gb18030_fallback_and_row_limit(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(settings, "max_parse_rows", 2)
    csv_bytes = "name,value\n中文,1\nextra,2\n".encode("gb18030")
    document = _upload("gb.csv", csv_bytes, "text/csv")

    response = _parse(document["id"])

    assert response.status_code == 200
    parsed_path = settings.data_path / response.json()["data"]["parsed_relative_path"]
    text = parsed_path.read_text(encoding="utf-8")
    assert "中文" in text
    assert "extra" not in text


def test_xlsx_parse_limits_sheets_and_rows(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    monkeypatch.setattr(settings, "max_parse_sheets", 1)
    monkeypatch.setattr(settings, "max_parse_rows", 1)
    workbook = Workbook()
    first = workbook.active
    first.title = "First"
    first.append(["name", "value"])
    first.append(["row1", "1"])
    first.append(["row2", "2"])
    second = workbook.create_sheet("Second")
    second.append(["hidden"])
    second.append(["nope"])
    path = tmp_path / "book.xlsx"
    workbook.save(path)
    document = _upload("book.xlsx", path.read_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    response = _parse(document["id"])

    assert response.status_code == 200
    text = (settings.data_path / response.json()["data"]["parsed_relative_path"]).read_text(encoding="utf-8")
    assert "First" in text
    assert "Second" not in text
    assert "row1" in text
    assert "row2" not in text


def test_docx_parse_success(tmp_path: Path) -> None:
    docx = DocxDocument()
    docx.add_paragraph("Docx text")
    path = tmp_path / "doc.docx"
    docx.save(path)
    document = _upload("doc.docx", path.read_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")

    response = _parse(document["id"])

    assert response.status_code == 200
    text = (settings.data_path / response.json()["data"]["parsed_relative_path"]).read_text(encoding="utf-8")
    assert "Docx text" in text


def test_pdf_parse_success(tmp_path: Path) -> None:
    path = tmp_path / "doc.pdf"
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "PDF text")
    pdf.save(path)
    pdf.close()
    document = _upload("doc.pdf", path.read_bytes(), "application/pdf")

    response = _parse(document["id"])

    assert response.status_code == 200
    text = (settings.data_path / response.json()["data"]["parsed_relative_path"]).read_text(encoding="utf-8")
    assert "PDF text" in text
